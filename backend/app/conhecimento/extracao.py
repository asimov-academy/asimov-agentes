"""Texto de um material da base: PDF, DOCX, TXT, MD ou uma página de site.

Só texto: imagem, áudio e vídeo não entram aqui. O que o modelo lê de mídia continua sendo do
`midia/`, por conversa, e nunca vira base de conhecimento.
"""

import html
import re
from io import BytesIO

import docx
from pypdf import PdfReader

LIMITE_DE_CARACTERES = 2_000_000
"""Cerca de 500 mil tokens. Acima disso o arquivo é grande demais para uma base de atendimento."""


class NaoDeuParaLer(ValueError):
    """Mensagem em português, pronta para a tela."""


TIPOS = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "texto",
    "text/markdown": "texto",
    "text/html": "html",
}

EXTENSOES = {".pdf": "pdf", ".docx": "docx", ".txt": "texto", ".md": "texto", ".markdown": "texto"}

_SCRIPT = re.compile(r"<(script|style|noscript)[^>]*>.*?</\1>", re.IGNORECASE | re.DOTALL)
_TAG = re.compile(r"<[^>]+>")
_ESPACO = re.compile(r"[ \t]+")
_LINHAS = re.compile(r"\n{3,}")


def formato(nome_arquivo: str, tipo_mime: str) -> str:
    """Pelo tipo declarado e, falhando, pela extensão: navegador manda `application/octet-stream`."""
    if tipo_mime in TIPOS:
        return TIPOS[tipo_mime]
    for extensao, qual in EXTENSOES.items():
        if nome_arquivo.lower().endswith(extensao):
            return qual
    raise NaoDeuParaLer("formato que não consigo ler; envie PDF, DOCX, TXT ou MD")


def de_pdf(conteudo: bytes) -> str:
    leitor = PdfReader(BytesIO(conteudo))
    return "\n\n".join((pagina.extract_text() or "").strip() for pagina in leitor.pages)


def de_docx(conteudo: bytes) -> str:
    arquivo = docx.Document(BytesIO(conteudo))
    partes = [p.text for p in arquivo.paragraphs]
    # Tabela de preço em .docx costuma ser tabela mesmo: sem isto, o documento chega vazio.
    for tabela in arquivo.tables:
        for linha in tabela.rows:
            celulas = [c.text.strip() for c in linha.cells if c.text.strip()]
            if celulas:
                partes.append(" | ".join(celulas))
    return "\n".join(partes)


def de_html(conteudo: bytes) -> str:
    texto = conteudo.decode("utf-8", errors="replace")
    texto = _SCRIPT.sub(" ", texto)
    texto = _TAG.sub("\n", texto)
    return html.unescape(texto)


def de_texto(conteudo: bytes) -> str:
    return conteudo.decode("utf-8", errors="replace")


LEITORES = {"pdf": de_pdf, "docx": de_docx, "html": de_html, "texto": de_texto}


def limpa(texto: str) -> str:
    texto = texto.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    texto = _ESPACO.sub(" ", texto)
    return _LINHAS.sub("\n\n", texto).strip()


def texto_de(conteudo: bytes, nome_arquivo: str, tipo_mime: str) -> str:
    """Levanta `NaoDeuParaLer` com o motivo: é o que vai para o status do documento e para a tela."""
    qual = formato(nome_arquivo, tipo_mime)
    try:
        bruto = LEITORES[qual](conteudo)
    except NaoDeuParaLer:
        raise
    except Exception as erro:  # o arquivo pode estar corrompido, protegido por senha ou vazio
        raise NaoDeuParaLer("não consegui abrir o arquivo; ele pode estar protegido ou corrompido") from erro
    texto = limpa(bruto)
    if not texto:
        raise NaoDeuParaLer("o arquivo não tem texto; PDF que é só imagem não serve")
    if len(texto) > LIMITE_DE_CARACTERES:
        raise NaoDeuParaLer("o arquivo é grande demais; divida em partes menores")
    return texto
